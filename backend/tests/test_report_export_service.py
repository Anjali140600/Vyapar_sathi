from datetime import date
from decimal import Decimal
from types import SimpleNamespace
import unittest

from openpyxl import load_workbook

from backend.app.services.report_export_service import ReportExportService


class ReportExportServiceTests(unittest.TestCase):
    def transactions(self):
        common = {
            "date": date(2026, 9, 1),
            "description": "Test entry",
            "gst_amount": Decimal("18"),
            "amount_paid": Decimal("118"),
            "amount_due": Decimal("0"),
            "due_date": None,
        }
        return [
            SimpleNamespace(**common, type="sales", category="Sales", amount=Decimal("500")),
            SimpleNamespace(**common, type="purchase", category="Stock", amount=Decimal("118")),
        ]

    def test_summary_uses_income_and_expense_types(self):
        summary = ReportExportService.summary(self.transactions())
        self.assertEqual(summary["sales"], 500.0)
        self.assertEqual(summary["expenses"], 118.0)
        self.assertEqual(summary["profit"], 382.0)
        self.assertEqual(summary["gst"], 36.0)

    def test_excel_contains_summary_and_transactions(self):
        output = ReportExportService.build_excel(self.transactions(), "Test period")
        workbook = load_workbook(output, data_only=True)
        self.assertEqual(workbook.sheetnames, ["Summary", "Transactions"])
        self.assertEqual(workbook["Summary"]["B5"].value, 500)
        self.assertEqual(workbook["Transactions"].max_row, 3)

    def test_pdf_has_valid_signature(self):
        output = ReportExportService.build_pdf(self.transactions(), "Test period")
        self.assertTrue(output.read(5).startswith(b"%PDF"))


if __name__ == "__main__":
    unittest.main()
