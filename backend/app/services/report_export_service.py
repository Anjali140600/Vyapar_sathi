"""Generate authenticated transaction reports entirely in memory."""

from datetime import date
from io import BytesIO

from backend.app.services.data_service import DataService


class ReportExportService:
    HEADERS = [
        "Date", "Type", "Category", "Description", "Amount (INR)",
        "GST (INR)", "Paid (INR)", "Due (INR)", "Due date",
    ]

    @staticmethod
    def rows(transactions) -> list[list]:
        return [
            [
                str(tx.date or ""),
                tx.type or "",
                tx.category or "General",
                tx.description or "",
                float(tx.amount or 0),
                float(tx.gst_amount or 0),
                float(tx.amount_paid or 0),
                float(tx.amount_due or 0),
                str(tx.due_date or ""),
            ]
            for tx in transactions
        ]

    @staticmethod
    def summary(transactions) -> dict:
        income = sum(
            float(tx.amount or 0)
            for tx in transactions
            if tx.type in DataService.INCOME_TYPES
        )
        expenses = sum(
            float(tx.amount or 0)
            for tx in transactions
            if tx.type in DataService.EXPENSE_TYPES
        )
        gst = sum(float(tx.gst_amount or 0) for tx in transactions)
        return {
            "sales": round(income, 2),
            "expenses": round(expenses, 2),
            "profit": round(income - expenses, 2),
            "gst": round(gst, 2),
        }

    @classmethod
    def build_excel(cls, transactions, period_label: str) -> BytesIO:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        summary_sheet.append(["Vyapar Sathi Business Report"])
        summary_sheet.append(["Period", period_label])
        summary_sheet.append(["Generated", date.today().isoformat()])
        summary_sheet.append([])
        summary = cls.summary(transactions)
        for label, key in (
            ("Total sales", "sales"),
            ("Total expenses", "expenses"),
            ("Net profit", "profit"),
            ("GST tracked", "gst"),
        ):
            summary_sheet.append([label, summary[key]])
        summary_sheet["A1"].font = Font(size=16, bold=True)
        summary_sheet.column_dimensions["A"].width = 24
        summary_sheet.column_dimensions["B"].width = 18
        for row in range(5, 9):
            summary_sheet.cell(row, 2).number_format = '₹#,##0.00'

        transaction_sheet = workbook.create_sheet("Transactions")
        transaction_sheet.append(cls.HEADERS)
        for row in cls.rows(transactions):
            transaction_sheet.append(row)
        header_fill = PatternFill("solid", fgColor="0F172A")
        for cell in transaction_sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for column in (5, 6, 7, 8):
            for row in range(2, transaction_sheet.max_row + 1):
                transaction_sheet.cell(row, column).number_format = '₹#,##0.00'
        widths = [13, 16, 22, 36, 16, 14, 14, 14, 13]
        for index, width in enumerate(widths, start=1):
            transaction_sheet.column_dimensions[
                transaction_sheet.cell(1, index).column_letter
            ].width = width
        transaction_sheet.freeze_panes = "A2"
        transaction_sheet.auto_filter.ref = transaction_sheet.dimensions

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    @classmethod
    def build_pdf(cls, transactions, period_label: str) -> BytesIO:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        output = BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=10 * mm,
            leftMargin=10 * mm,
            topMargin=10 * mm,
            bottomMargin=10 * mm,
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "ReportTitle", parent=styles["Title"], alignment=TA_CENTER, fontSize=18
        )
        summary = cls.summary(transactions)
        story = [
            Paragraph("Vyapar Sathi Business Report", title_style),
            Paragraph(f"Period: {period_label} | Generated: {date.today().isoformat()}", styles["Normal"]),
            Spacer(1, 5 * mm),
        ]
        summary_data = [
            ["Sales", "Expenses", "Net profit", "GST tracked"],
            [
                f"INR {summary['sales']:,.2f}",
                f"INR {summary['expenses']:,.2f}",
                f"INR {summary['profit']:,.2f}",
                f"INR {summary['gst']:,.2f}",
            ],
        ]
        summary_table = Table(summary_data, colWidths=[65 * mm] * 4)
        summary_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("PADDING", (0, 0), (-1, -1), 7),
        ]))
        story.extend([summary_table, Spacer(1, 6 * mm)])

        table_data = [cls.HEADERS] + [
            [
                *row[:4],
                *[f"{value:,.2f}" for value in row[4:8]],
                row[8],
            ]
            for row in cls.rows(transactions)
        ]
        if len(table_data) == 1:
            table_data.append(["No transactions found", "", "", "", "", "", "", "", ""])
        transaction_table = Table(
            table_data,
            repeatRows=1,
            colWidths=[20*mm, 23*mm, 30*mm, 53*mm, 25*mm, 22*mm, 22*mm, 22*mm, 20*mm],
        )
        transaction_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0EA5E9")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("PADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(transaction_table)
        document.build(story)
        output.seek(0)
        return output
